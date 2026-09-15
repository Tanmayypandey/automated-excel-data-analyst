"""Optimized Excel Exporter Module.

Creates a professional Excel report while remaining efficient for
medium and large datasets.

Sheets:
1. Dashboard
2. Cleaned Data
3. Cleaning Report
4. Original Data
5. Chart Data
"""

import os
from typing import Dict, Any, List

import pandas as pd
import numpy as np
import openpyxl

from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.cleaning_report_generator import populate_cleaning_report
from src.dashboard_generator import build_dashboard


# ---------------------------------------------------------
# EXPORT LIMITS
# ---------------------------------------------------------

# Full cleaned dataset can be exported up to this limit.
MAX_CLEANED_EXPORT_ROWS = 100_000

# Raw data is useful for comparison, but duplicating 100k+
# rows makes the workbook unnecessarily heavy.
FULL_RAW_EXPORT_THRESHOLD = 50_000
RAW_PREVIEW_ROWS = 5_000


def export_excel_report(
    raw_df: pd.DataFrame,
    cleaned_df: pd.DataFrame,
    cleaning_stats: Dict[str, Any],
    kpis: List[Dict[str, Any]],
    chart_specs: List[Dict[str, Any]],
    output_file_path: str
) -> str:
    """Generate optimized Excel analysis workbook."""

    output_file_path = os.path.abspath(output_file_path)

    os.makedirs(
        os.path.dirname(output_file_path),
        exist_ok=True
    )

    wb = openpyxl.Workbook()

    # -----------------------------------------------------
    # CREATE SHEETS
    # -----------------------------------------------------

    ws_dashboard = wb.active
    ws_dashboard.title = "Dashboard"

    ws_cleaned = wb.create_sheet("Cleaned Data")
    ws_report = wb.create_sheet("Cleaning Report")
    ws_original = wb.create_sheet("Original Data")
    ws_chart_data = wb.create_sheet("Chart Data")

    # -----------------------------------------------------
    # ORIGINAL DATA
    # -----------------------------------------------------

    if len(raw_df) <= FULL_RAW_EXPORT_THRESHOLD:

        raw_export_df = raw_df

        original_title = (
            f"Original Data — {len(raw_df):,} rows"
        )

    else:

        raw_export_df = raw_df.head(RAW_PREVIEW_ROWS)

        original_title = (
            f"Original Data Preview — first "
            f"{RAW_PREVIEW_ROWS:,} of {len(raw_df):,} rows"
        )

    _write_dataframe_fast(
        ws_original,
        raw_export_df,
        original_title
    )

    # Free reference as soon as possible.
    del raw_export_df

    # -----------------------------------------------------
    # CLEANED DATA
    # -----------------------------------------------------

    cleaned_rows_to_export = min(
        len(cleaned_df),
        MAX_CLEANED_EXPORT_ROWS
    )

    cleaned_export_df = cleaned_df.iloc[
        :cleaned_rows_to_export
    ]

    if len(cleaned_df) > MAX_CLEANED_EXPORT_ROWS:

        cleaned_title = (
            f"Cleaned Data — first "
            f"{MAX_CLEANED_EXPORT_ROWS:,} of "
            f"{len(cleaned_df):,} rows"
        )

    else:

        cleaned_title = (
            f"Cleaned Data — {len(cleaned_df):,} rows"
        )

    _write_dataframe_fast(
        ws_cleaned,
        cleaned_export_df,
        cleaned_title
    )

    del cleaned_export_df

    # -----------------------------------------------------
    # CLEANING REPORT
    # -----------------------------------------------------

    populate_cleaning_report(
        ws_report,
        cleaning_stats
    )

    # -----------------------------------------------------
    # CHART DATA
    # -----------------------------------------------------

    chart_ranges = _populate_chart_data(
        ws_chart_data,
        cleaned_df,
        chart_specs
    )

    # -----------------------------------------------------
    # DATASET SUMMARY
    # -----------------------------------------------------

    dataset_summary = {
        "total_rows": len(cleaned_df),
        "total_columns": len(cleaned_df.columns),

        "duplicate_rows":
            cleaning_stats.get(
                "duplicates_removed",
                0
            ),

        "missing_values_cleaned": max(
            0,
            cleaning_stats.get(
                "missing_values_before",
                0
            )
            -
            cleaning_stats.get(
                "missing_values_after",
                0
            )
        )
    }

    # -----------------------------------------------------
    # DASHBOARD
    # -----------------------------------------------------

    build_dashboard(
        ws_dash=ws_dashboard,
        ws_chart_data=ws_chart_data,
        kpis=kpis,
        chart_specs=chart_specs,
        chart_data_ranges=chart_ranges,
        dataset_summary=dataset_summary
    )

    # Hide chart data because end users normally
    # don't need to see chart aggregation tables.
    ws_chart_data.sheet_state = "hidden"

    # -----------------------------------------------------
    # SAVE
    # -----------------------------------------------------

    wb.save(output_file_path)
    wb.close()

    return output_file_path


# =========================================================
# FAST DATAFRAME WRITER
# =========================================================

def _excel_safe_value(value):
    """Convert pandas/numpy values into Excel-safe Python values."""

    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    if isinstance(value, np.datetime64):
        try:
            return pd.Timestamp(value).to_pydatetime()
        except Exception:
            return str(value)

    if isinstance(value, np.integer):
        return int(value)

    if isinstance(value, np.floating):
        return float(value)

    if isinstance(value, np.bool_):
        return bool(value)

    # Excel cannot store timezone-aware datetimes.
    if hasattr(value, "tzinfo") and value.tzinfo is not None:
        return str(value)

    return value


def _write_dataframe_fast(
    ws: Worksheet,
    df: pd.DataFrame,
    title: str
) -> None:
    """Write DataFrame efficiently.

    Important performance improvement:
    Data rows are NOT individually styled.

    Only title/header rows receive styling.
    """

    ws.sheet_view.showGridLines = True

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    number_of_columns = max(
        len(df.columns),
        1
    )

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=number_of_columns
    )

    title_cell = ws.cell(
        row=1,
        column=1,
        value=title
    )

    title_cell.font = Font(
        name="Segoe UI",
        size=11,
        bold=True
    )

    title_cell.alignment = Alignment(
        vertical="center"
    )

    # -----------------------------------------------------
    # HEADER
    # -----------------------------------------------------

    header_fill = PatternFill(
        start_color="1E293B",
        end_color="1E293B",
        fill_type="solid"
    )

    header_font = Font(
        name="Segoe UI",
        size=10,
        bold=True,
        color="FFFFFF"
    )

    for column_index, column_name in enumerate(
        df.columns,
        start=1
    ):

        cell = ws.cell(
            row=2,
            column=column_index,
            value=str(column_name)
        )

        cell.font = header_font
        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # -----------------------------------------------------
    # DATA
    # -----------------------------------------------------

    # ws.append is significantly lighter than constructing
    # and styling every individual cell.
    for row in df.itertuples(
        index=False,
        name=None
    ):

        ws.append([
            _excel_safe_value(value)
            for value in row
        ])

    # -----------------------------------------------------
    # WORKSHEET SETTINGS
    # -----------------------------------------------------

    ws.freeze_panes = "A3"

    if len(df.columns) > 0:
        ws.auto_filter.ref = (
            f"A2:"
            f"{get_column_letter(len(df.columns))}"
            f"{len(df) + 2}"
        )

    # Width is based only on column names.
    # Do NOT scan entire columns for width calculation.
    for column_index, column_name in enumerate(
        df.columns,
        start=1
    ):

        width = min(
            max(
                len(str(column_name)) + 4,
                12
            ),
            30
        )

        ws.column_dimensions[
            get_column_letter(column_index)
        ].width = width


# =========================================================
# CHART DATA
# =========================================================

def _populate_chart_data(
    ws: Worksheet,
    df: pd.DataFrame,
    chart_specs: List[Dict[str, Any]]
) -> List[Dict[str, Any]]:
    """Create small aggregated tables used by dashboard charts."""

    ws.sheet_view.showGridLines = True

    header_fill = PatternFill(
        start_color="3B82F6",
        end_color="3B82F6",
        fill_type="solid"
    )

    header_font = Font(
        name="Segoe UI",
        size=9,
        bold=True,
        color="FFFFFF"
    )

    ranges: List[Dict[str, Any]] = []

    column_offsets = [
        (1, 2),
        (4, 5),
        (7, 8),
        (10, 11)
    ]

    for index, spec in enumerate(
        chart_specs[:4]
    ):

        category_column, value_column = (
            column_offsets[index]
        )

        chart_type = spec.get(
            "chart_type",
            "bar"
        )

        x_col = spec.get("x_col")
        y_col = spec.get("y_col")

        aggregation = spec.get(
            "aggregation",
            "sum"
        )

        table_df = pd.DataFrame()

        try:

            # ---------------------------------------------
            # LINE
            # ---------------------------------------------

            if (
                chart_type == "line"
                and x_col in df.columns
                and y_col in df.columns
            ):

                temp = df[
                    [x_col, y_col]
                ].dropna()

                temp = temp.copy()

                temp[x_col] = pd.to_datetime(
                    temp[x_col],
                    errors="coerce"
                )

                temp = temp.dropna(
                    subset=[x_col]
                )

                # Aggregate by day.
                temp["_chart_date"] = (
                    temp[x_col]
                    .dt.strftime("%Y-%m-%d")
                )

                if aggregation == "mean":

                    table_df = (
                        temp.groupby(
                            "_chart_date",
                            as_index=False
                        )[y_col]
                        .mean()
                    )

                else:

                    table_df = (
                        temp.groupby(
                            "_chart_date",
                            as_index=False
                        )[y_col]
                        .sum()
                    )

                table_df = (
                    table_df
                    .sort_values("_chart_date")
                    .tail(30)
                )

            # ---------------------------------------------
            # BAR
            # ---------------------------------------------

            elif (
                chart_type == "bar"
                and x_col in df.columns
                and y_col in df.columns
            ):

                temp = df[
                    [x_col, y_col]
                ].dropna()

                if aggregation == "mean":

                    table_df = (
                        temp.groupby(
                            x_col,
                            as_index=False
                        )[y_col]
                        .mean()
                    )

                else:

                    table_df = (
                        temp.groupby(
                            x_col,
                            as_index=False
                        )[y_col]
                        .sum()
                    )

                limit = (
                    10
                    if "top10" in aggregation
                    else 20
                )

                table_df = (
                    table_df
                    .sort_values(
                        y_col,
                        ascending=False
                    )
                    .head(limit)
                )

            # ---------------------------------------------
            # PIE
            # ---------------------------------------------

            elif (
                chart_type == "pie"
                and x_col in df.columns
                and y_col in df.columns
            ):

                temp = df[
                    [x_col, y_col]
                ].dropna()

                table_df = (
                    temp.groupby(
                        x_col,
                        as_index=False
                    )[y_col]
                    .sum()
                    .sort_values(
                        y_col,
                        ascending=False
                    )
                    .head(6)
                )

            # ---------------------------------------------
            # SCATTER
            # ---------------------------------------------

            elif (
                chart_type == "scatter"
                and x_col in df.columns
                and y_col in df.columns
            ):

                temp = df[
                    [x_col, y_col]
                ].dropna()

                if len(temp) > 100:

                    temp = temp.sample(
                        n=100,
                        random_state=42
                    )

                table_df = temp.sort_values(
                    x_col
                )

            # ---------------------------------------------
            # HISTOGRAM
            # ---------------------------------------------

            elif (
                chart_type == "histogram"
                and x_col in df.columns
            ):

                series = pd.to_numeric(
                    df[x_col],
                    errors="coerce"
                ).dropna()

                if len(series) > 0:

                    counts, bin_edges = np.histogram(
                        series,
                        bins=8
                    )

                    labels = [
                        f"{bin_edges[i]:.1f}-"
                        f"{bin_edges[i + 1]:.1f}"

                        for i in range(
                            len(counts)
                        )
                    ]

                    table_df = pd.DataFrame({
                        "Interval": labels,
                        "Frequency": counts
                    })

        except Exception:

            table_df = pd.DataFrame({
                "Category": ["No Data"],
                "Value": [0]
            })

        # ---------------------------------------------
        # FALLBACK
        # ---------------------------------------------

        if table_df.empty:

            table_df = pd.DataFrame({
                "Category": ["No Data"],
                "Value": [0]
            })

        # Excel dashboard expects two columns.
        if len(table_df.columns) < 2:

            table_df["Value"] = 0

        table_df = table_df.iloc[:, :2]

        # ---------------------------------------------
        # WRITE CHART TABLE
        # ---------------------------------------------

        headers = list(
            table_df.columns
        )

        first_header = ws.cell(
            row=1,
            column=category_column,
            value=str(headers[0])
        )

        first_header.font = header_font
        first_header.fill = header_fill

        second_header = ws.cell(
            row=1,
            column=value_column,
            value=str(headers[1])
        )

        second_header.font = header_font
        second_header.fill = header_fill

        for row_offset, row in enumerate(
            table_df.itertuples(
                index=False,
                name=None
            ),
            start=2
        ):

            category_value = row[0]
            numeric_value = row[1]

            ws.cell(
                row=row_offset,
                column=category_column,
                value=str(category_value)
            )

            ws.cell(
                row=row_offset,
                column=value_column,
                value=_excel_safe_value(
                    numeric_value
                )
            )

        end_row = len(table_df) + 1

        ranges.append({
            "cat_col": category_column,
            "val_col": value_column,
            "start_row": 1,
            "end_row": end_row
        })

        ws.column_dimensions[
            get_column_letter(
                category_column
            )
        ].width = 18

        ws.column_dimensions[
            get_column_letter(
                value_column
            )
        ].width = 14

    return ranges