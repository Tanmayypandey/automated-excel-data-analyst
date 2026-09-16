"""Lightweight Output Validator Module.

Designed for large generated Excel workbooks.

Validates:
- File exists
- File is not empty
- File is a valid XLSX/ZIP container
- Required worksheets exist
- Data sheets contain rows

Important:
The workbook is opened in read_only mode to avoid loading
large Excel files fully into memory.
"""

import os
import zipfile
from typing import Dict, Any

import openpyxl


REQUIRED_SHEETS = [
    "Dashboard",
    "Cleaned Data",
    "Cleaning Report",
    "Original Data",
    "Chart Data",
]


def _error_result(
    error: str,
    file_size: int = 0,
    sheet_names=None,
    original_rows: int = 0,
    cleaned_rows: int = 0,
) -> Dict[str, Any]:
    """Create a consistent validation error response."""

    return {
        "valid": False,
        "error": error,
        "file_size": file_size,
        "sheet_names": sheet_names or [],
        "chart_count": 0,
        "original_rows": original_rows,
        "cleaned_rows": cleaned_rows,
    }


def validate_output(output_path: str) -> Dict[str, Any]:
    """Validate a generated Excel workbook using minimal memory."""

    # ---------------------------------------------------------
    # 1. Check file existence
    # ---------------------------------------------------------
    if not os.path.exists(output_path):
        return _error_result(
            f"Output file does not exist: {output_path}"
        )

    # ---------------------------------------------------------
    # 2. Check file size
    # ---------------------------------------------------------
    try:
        size = os.path.getsize(output_path)
    except OSError as exc:
        return _error_result(
            f"Unable to read output file information: {exc}"
        )

    if size <= 0:
        return _error_result(
            "Generated output file is empty (0 bytes)."
        )

    # ---------------------------------------------------------
    # 3. Check XLSX container integrity
    #
    # XLSX files are ZIP containers. This gives us a cheap
    # corruption check without loading all worksheet cells.
    # ---------------------------------------------------------
    if not zipfile.is_zipfile(output_path):
        return _error_result(
            "Generated output is not a valid XLSX file.",
            file_size=size,
        )

    try:
        with zipfile.ZipFile(output_path, "r") as archive:
            bad_file = archive.testzip()

            if bad_file is not None:
                return _error_result(
                    f"Generated XLSX contains a corrupted component: {bad_file}",
                    file_size=size,
                )

    except Exception as exc:
        return _error_result(
            f"Failed XLSX integrity check: {exc}",
            file_size=size,
        )

    # ---------------------------------------------------------
    # 4. Open workbook in READ-ONLY mode
    #
    # This is the important optimization.
    # ---------------------------------------------------------
    wb = None

    try:
        wb = openpyxl.load_workbook(
            output_path,
            read_only=True,
            data_only=True,
        )

        sheet_names = list(wb.sheetnames)

        # -----------------------------------------------------
        # 5. Validate required sheets
        # -----------------------------------------------------
        missing_sheets = [
            sheet
            for sheet in REQUIRED_SHEETS
            if sheet not in sheet_names
        ]

        if missing_sheets:
            return _error_result(
                f"Missing required sheets: {', '.join(missing_sheets)}",
                file_size=size,
                sheet_names=sheet_names,
            )

        # -----------------------------------------------------
        # 6. Validate data sheets
        # -----------------------------------------------------
        ws_orig = wb["Original Data"]
        ws_clean = wb["Cleaned Data"]

        orig_rows = ws_orig.max_row or 0
        clean_rows = ws_clean.max_row or 0

        if orig_rows <= 1:
            return _error_result(
                "Original Data worksheet has no data rows.",
                file_size=size,
                sheet_names=sheet_names,
                original_rows=orig_rows,
                cleaned_rows=clean_rows,
            )

        if clean_rows <= 1:
            return _error_result(
                "Cleaned Data worksheet has no data rows.",
                file_size=size,
                sheet_names=sheet_names,
                original_rows=orig_rows,
                cleaned_rows=clean_rows,
            )

        # -----------------------------------------------------
        # 7. Success
        #
        # Chart validation is intentionally skipped here.
        # Reading chart objects requires a heavier workbook
        # load and is unnecessary after successful generation.
        # -----------------------------------------------------
        return {
            "valid": True,
            "error": None,
            "file_size": size,
            "sheet_names": sheet_names,
            "chart_count": None,
            "original_rows": max(orig_rows - 1, 0),
            "cleaned_rows": max(clean_rows - 1, 0),
        }

    except Exception as exc:
        return _error_result(
            f"Failed to validate workbook: {exc}",
            file_size=size,
        )

    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass